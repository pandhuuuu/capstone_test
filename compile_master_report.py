import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

base_dir = os.path.dirname(os.path.abspath(__file__))
docs_dir = os.path.join(base_dir, 'docs', 'capstone')
diagrams_dir = os.path.join(docs_dir, 'diagrams')
master_md = os.path.join(docs_dir, 'LAPORAN_LENGKAP_CAPSTONE_BNPB.md')
output_docx = os.path.join(docs_dir, 'Laporan_Akhir_Capstone_Pusdalops_BNPB.docx')

doc = Document()

# Margin Standar Institusi FRI Telkom: 4 cm kiri, 3 cm atas, kanan, bawah
for section in doc.sections:
    section.top_margin = Inches(1.18)    # ~3 cm
    section.bottom_margin = Inches(1.18) # ~3 cm
    section.left_margin = Inches(1.57)   # ~4 cm
    section.right_margin = Inches(1.18)  # ~3 cm

# Base style
style_normal = doc.styles['Normal']
style_normal.font.name = 'Times New Roman'
style_normal.font.size = Pt(12)
style_normal.font.color.rgb = RGBColor(0, 0, 0)
style_normal.paragraph_format.line_spacing = 1.5
style_normal.paragraph_format.space_after = Pt(6)

def add_formatted_runs(p, text, base_font="Times New Roman", base_size=Pt(12), base_bold=False, base_italic=False):
    """
    Parser cerdas untuk mengubah sintaks markdown (**bold**, *italic*, ***bold-italic***, `code`)
    menjadi Word runs berformat tanpa meninggalkan satupun tanda bintang (*) atau backtick di naskah.
    """
    pattern = re.compile(r'(\*\*\*.*?\*\*\*|\*\*.*?\*\*|\*[^*\n]+?\*|`.*?`)')
    tokens = pattern.split(text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('***') and token.endswith('***') and len(token) >= 6:
            clean_token = token[3:-3].replace('*', '')
            r = p.add_run(clean_token)
            r.bold = True
            r.italic = True
        elif token.startswith('**') and token.endswith('**') and len(token) >= 4:
            inner = token[2:-2]
            if '(*' in inner and '*)' in inner:
                sub_parts = re.split(r'(\(\*.*?\*\))', inner)
                for sp in sub_parts:
                    if sp.startswith('(*') and sp.endswith('*)'):
                        sub_r = p.add_run(f"({sp[2:-2].replace('*', '')})")
                        sub_r.bold = True
                        sub_r.italic = True
                        sub_r.font.name = base_font
                        sub_r.font.size = base_size
                    else:
                        clean_sp = sp.replace('*', '')
                        sub_r = p.add_run(clean_sp)
                        sub_r.bold = True
                        sub_r.font.name = base_font
                        sub_r.font.size = base_size
                continue
            else:
                clean_inner = inner.replace('*', '')
                r = p.add_run(clean_inner)
                r.bold = True
        elif token.startswith('*') and token.endswith('*') and len(token) >= 2:
            clean_token = token[1:-1].replace('*', '')
            r = p.add_run(clean_token)
            r.italic = True
        elif token.startswith('`') and token.endswith('`') and len(token) >= 2:
            r = p.add_run(token[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(base_size.pt - 1)
            continue
        else:
            clean_token = token.replace('*', '')
            r = p.add_run(clean_token)
            if base_bold:
                r.bold = True
            if base_italic:
                r.italic = True
                
        r.font.name = base_font
        r.font.size = base_size

def add_heading_1(text):
    p = doc.add_paragraph()
    is_bab = text.startswith('BAB ') or text.startswith('LAPORAN ') or text in ['ABSTRAK', 'ABSTRACT', 'DAFTAR PUSTAKA', 'DAFTAR ISTILAH DAN SINGKATAN']
    if is_bab:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    add_formatted_runs(p, text, base_font="Times New Roman", base_size=Pt(14), base_bold=True)

def add_heading_2(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    add_formatted_runs(p, text, base_font="Times New Roman", base_size=Pt(12.5), base_bold=True)

def add_heading_3(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    add_formatted_runs(p, text, base_font="Times New Roman", base_size=Pt(12), base_bold=True)

def add_heading_4(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    add_formatted_runs(p, text, base_font="Times New Roman", base_size=Pt(12), base_bold=True, base_italic=True)

def add_body_paragraph(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(6)
    add_formatted_runs(p, text, base_font="Times New Roman", base_size=Pt(12))

def add_bullet(text):
    p = doc.add_paragraph(style='List Bullet')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(3)
    add_formatted_runs(p, text, base_font="Times New Roman", base_size=Pt(12))

def add_diagram_image(image_filename, caption_text):
    img_path = os.path.join(diagrams_dir, image_filename)
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(10)
        p_img.paragraph_format.space_after = Pt(4)
        run = p_img.add_run()
        # 5.5 inci tepat pada batas margin cetak A4 (8.27 - 1.57 - 1.18 = 5.52 inci)
        run.add_picture(img_path, width=Inches(5.5))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        # Caption standar FRI Telkom: Times New Roman 10pt italic di bawah gambar
        clean_cap = caption_text.replace('*', '')
        add_formatted_runs(p_cap, clean_cap, base_font="Times New Roman", base_size=Pt(10), base_italic=True)
        p_cap.paragraph_format.space_after = Pt(12)

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

# Mapping nama gambar berdasarkan teks caption
diagram_map = {
    'Gambar 1.1': ('gambar_1_2_fishbone_diagram.png', 'Gambar 1.1 Diagram Tulang Ikan (Fishbone Diagram) Analisis Akar Masalah Meja Piket Pusdalops BNPB'),
    'Gambar 1.2': ('gambar_bpmn_asis.png', 'Gambar 1.2 Business Process Model and Notation (BPMN 2.0) Alur Operasional Eksisting (As-Is) Meja Piket Pusdalops BNPB (Abid)'),
    'Gambar 1.3': ('gambar_bpmn_tobe.png', 'Gambar 1.3 Business Process Model and Notation (BPMN 2.0) Alur Operasional Usulan (To-Be) Sistem Otomasi Terintegrasi (Abid)'),
    'Gambar 3.1': ('gambar_3_1_sistematika_perancangan.png', 'Gambar 3.1 Diagram Alir Sistematika Perancangan Sistem Informasi Pusdalops BNPB'),
    'Gambar 4.1': ('gambar_usecase_diagram.png', 'Gambar 4.1 Use Case Diagram Sistem Informasi Pusdalops BNPB (Adel)'),
    'Gambar 4.2': ('gambar_activity_diagram.png', 'Gambar 4.2 Activity Diagram Verifikasi, Deduplikasi, dan Ekspor Laporan Pusdalops BNPB (Abid)'),
    'Gambar 4.3': ('gambar_sequence_diagram.png', 'Gambar 4.3 Sequence Diagram Skenario Ingestion, Verifikasi 5 Blok, dan Ekspor Lapsit Word (Pandhu)'),
    'Gambar 4.4': ('gambar_class_diagram.png', 'Gambar 4.4 Class Diagram Struktur Arsitektur Sistem Informasi Pusdalops BNPB (Adel)'),
    'Gambar 4.5': ('gambar_erd_database.png', 'Gambar 4.5 Entity Relationship Diagram (ERD) Normalisasi Bentuk Ketiga / 3NF (Seto)'),
}

with open(master_md, 'r', encoding='utf-8') as f:
    lines = f.readlines()

in_mermaid = False
in_code = False
in_table = False
table_lines = []

for line in lines:
    stripped = line.strip()
    
    # Lewati blok kode mermaid karena sudah digantikan oleh gambar diagram HD
    if stripped.startswith('```mermaid'):
        in_mermaid = True
        continue
    elif in_mermaid and stripped.startswith('```'):
        in_mermaid = False
        continue
    elif in_mermaid:
        continue

    # Lewati blok kode umum
    if stripped.startswith('```'):
        in_code = not in_code
        continue
    if in_code:
        continue

    # Cek penyisipan diagram gambar berdasarkan label Gambar X.X
    matched_img = False
    for g_key, (img_fn, cap_txt) in diagram_map.items():
        if stripped.startswith(g_key):
            add_diagram_image(img_fn, stripped)
            matched_img = True
            break
    if matched_img:
        continue

    # Deteksi Tabel Markdown
    if stripped.startswith('|') and stripped.endswith('|'):
        table_lines.append(stripped)
        in_table = True
        continue
    else:
        if in_table and table_lines:
            rows_data = []
            for tline in table_lines:
                if '---' in tline:
                    continue
                cols = [c.strip() for c in tline.split('|')[1:-1]]
                if any(cols):
                    rows_data.append(cols)
            if rows_data:
                n_rows = len(rows_data)
                n_cols = len(rows_data[0])
                table = doc.add_table(rows=n_rows, cols=n_cols)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                table.autofit = True
                for r_idx, row in enumerate(rows_data):
                    for c_idx, val in enumerate(row):
                        if c_idx < n_cols:
                            cell = table.cell(r_idx, c_idx)
                            # Bersihkan teks cell bawaan python-docx
                            cell.text = ""
                            cp = cell.paragraphs[0]
                            cp.paragraph_format.line_spacing = 1.15
                            cp.paragraph_format.space_after = Pt(2)
                            cp.paragraph_format.space_before = Pt(2)
                            is_header_row = (r_idx == 0)
                            add_formatted_runs(cp, val, base_font="Times New Roman", base_size=Pt(9.5), base_bold=is_header_row)
                            if is_header_row:
                                set_cell_background(cell, "EAECEE")
                doc.add_paragraph().paragraph_format.space_after = Pt(4)
            table_lines = []
            in_table = False

    if not stripped:
        continue
        
    if stripped.startswith('# '):
        add_heading_1(stripped[2:])
    elif stripped.startswith('## '):
        add_heading_2(stripped[3:])
    elif stripped.startswith('### '):
        add_heading_3(stripped[4:])
    elif stripped.startswith('#### '):
        add_heading_4(stripped[5:])
    elif stripped.startswith('* ') or stripped.startswith('- '):
        add_bullet(stripped[2:])
    elif stripped.startswith('---'):
        doc.add_page_break()
    else:
        # Periksa apakah baris ini judul tabel (Tabel X.X ...)
        if stripped.startswith('Tabel '):
            p_tcap = doc.add_paragraph()
            p_tcap.alignment = WD_ALIGN_PARAGRAPH.LEFT
            clean_tcap = stripped.replace('*', '')
            add_formatted_runs(p_tcap, clean_tcap, base_font="Times New Roman", base_size=Pt(10.5), base_bold=True)
            p_tcap.paragraph_format.space_before = Pt(10)
            p_tcap.paragraph_format.space_after = Pt(2)
            p_tcap.paragraph_format.keep_with_next = True
        else:
            add_body_paragraph(stripped)

try:
    doc.save(output_docx)
    print(f"Master Docx berhasil dikompilasi secara sempurna ke: {output_docx}")
except PermissionError:
    alt_output_docx = os.path.join(docs_dir, 'Laporan_Akhir_Capstone_Pusdalops_BNPB_Revisi_Final.docx')
    doc.save(alt_output_docx)
    print(f"Pemberitahuan: File asli sedang dibuka di aplikasi Microsoft Word.")
    print(f"Master Docx berhasil dikompilasi ke versi baru yang bersih: {alt_output_docx}")
