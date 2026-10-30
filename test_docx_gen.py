from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import datetime

doc = Document()

# Set Standard Page Margins (Normal 1 inch)
for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

# Header BNPB
header = doc.add_heading(level=0)
run_header = header.add_run("BADAN NASIONAL PENANGGULANGAN BENCANA\nPUSAT PENGENDALIAN OPERASI (PUSDALOPS PB)")
run_header.font.name = "Arial"
run_header.font.size = Pt(13)
run_header.font.bold = True
run_header.font.color.rgb = RGBColor(0, 32, 96) # Dark Blue
header.alignment = WD_ALIGN_PARAGRAPH.CENTER

sub_header = doc.add_paragraph()
today_str = datetime.datetime.now().strftime("%d %B %Y").upper()
run_sub = sub_header.add_run(f"LAPORAN HARIAN SITUASI BENCANA\nPERIODE: {today_str}")
run_sub.font.name = "Arial"
run_sub.font.size = Pt(11)
run_sub.font.bold = True
sub_header.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Bagian A Heading
p_a = doc.add_paragraph()
run_a = p_a.add_run("BAGIAN A. RINCIAN KEJADIAN BENCANA")
run_a.font.name = "Arial"
run_a.font.size = Pt(11)
run_a.font.bold = True

# Sample Event 1
events = [
    {
        "no": "1",
        "title": "BANJIR - KABUPATEN BANDUNG, PROVINSI JAWA BARAT",
        "waktu": "Jumat, 09 Oktober 2026, Pkl. 14.30 WIB",
        "lokasi": "Kec. Dayeuhkolot, Kec. Baleendah, Kec. Bojongsoang",
        "kronologi": "Hujan dengan intensitas tinggi menyebabkan luapan Sungai Citarum merendam pemukiman warga.",
        "korban": "• Korban Jiwa: Nihil\n• Korban Mengungsi: 142 KK / 480 Jiwa\n• Terdampak: 1.250 KK",
        "kerusakan": "• 1.120 Unit Rumah terendam (TMA 50 - 120 cm)\n• 3 Unit Tempat Ibadah terdampak\n• 2 Unit Sarana Pendidikan terdampak",
        "upaya": "• BPBD Kab. Bandung melakukan asesmen dan evakuasi warga menggunakan perahu karet.\n• Pendirian posko pengungsian di Gedung Inkanas Dayeuhkolot.\n• Pendistribusian bantuan logistik makanan siap saji.",
        "kondisi": "Air berangsur surut 10-20 cm, namun cuaca terpantau masih mendung berpotensi hujan.",
        "sumber": "Pusdalops BPBD Kab. Bandung, Antara News, BMKG"
    }
]

for ev in events:
    p_title = doc.add_paragraph()
    r_t = p_title.add_run(f"{ev['no']}. {ev['title']}")
    r_t.font.name = "Arial"
    r_t.font.size = Pt(10)
    r_t.font.bold = True

    # Atribut Resmi BNPB Bagian A
    items = [
        ("Waktu Kejadian", ev["waktu"]),
        ("Lokasi", ev["lokasi"]),
        ("Kronologi", ev["kronologi"]),
        ("Korban Jiwa & Terdampak", ev["korban"]),
        ("Kerusakan Materil", ev["kerusakan"]),
        ("Upaya Penanganan", ev["upaya"]),
        ("Kondisi Mutakhir", ev["kondisi"]),
        ("Sumber Rujukan", ev["sumber"]),
    ]

    for label, val in items:
        p_item = doc.add_paragraph()
        p_item.paragraph_format.left_indent = Inches(0.25)
        p_item.paragraph_format.space_after = Pt(2)
        
        r_bullet = p_item.add_run(f"➢ {label}: ")
        r_bullet.font.name = "Arial"
        r_bullet.font.size = Pt(9.5)
        r_bullet.font.bold = True
        
        # Pisahkan baris jika ada newline
        lines = val.split("\n")
        r_val_first = p_item.add_run(lines[0])
        r_val_first.font.name = "Arial"
        r_val_first.font.size = Pt(9.5)
        
        for extra_line in lines[1:]:
            p_sub = doc.add_paragraph()
            p_sub.paragraph_format.left_indent = Inches(0.5)
            p_sub.paragraph_format.space_after = Pt(1)
            r_extra = p_sub.add_run(extra_line)
            r_extra.font.name = "Arial"
            r_extra.font.size = Pt(9.5)

output_path = "output_laporan_bagian_a_test.docx"
doc.save(output_path)
print(f"File Berhasil Dibuat: {output_path}")
