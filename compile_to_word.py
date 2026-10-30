import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

base_dir = os.path.dirname(os.path.abspath(__file__))
docs_dir = os.path.join(base_dir, 'docs', 'capstone')
diagrams_dir = os.path.join(docs_dir, 'diagrams')
output_docx = os.path.join(docs_dir, 'Laporan_Akhir_Capstone_Pusdalops_BNPB.docx')

doc = Document()

# Set margin 4-3-3-3 cm standar akademik
for section in doc.sections:
    section.top_margin = Inches(1.18)    # ~3 cm
    section.bottom_margin = Inches(1.18) # ~3 cm
    section.left_margin = Inches(1.57)   # ~4 cm
    section.right_margin = Inches(1.18)  # ~3 cm

# Set base style font ke Times New Roman
style_normal = doc.styles['Normal']
style_normal.font.name = 'Times New Roman'
style_normal.font.size = Pt(12)
style_normal.font.color.rgb = RGBColor(0, 0, 0)
style_normal.paragraph_format.line_spacing = 1.5
style_normal.paragraph_format.space_after = Pt(6)

def add_title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(0, 0, 0)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.space_before = Pt(12)

def add_heading_1(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(0, 0, 0)
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)

def add_heading_2(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12.5)
    r.font.color.rgb = RGBColor(0, 0, 0)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)

def add_heading_3(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(0, 0, 0)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)

def add_body_paragraph(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    # Parsing bold markdown **text**
    parts = re.split(r'(\*\*.*?\*\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            r = p.add_run(part[2:-2])
            r.bold = True
        elif part.startswith('*') and part.endswith('*') and len(part) > 2:
            r = p.add_run(part[1:-1])
            r.italic = True
        else:
            p.add_run(part)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(6)

def add_bullet(text):
    p = doc.add_paragraph(style='List Bullet')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    parts = re.split(r'(\*\*.*?\*\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            r = p.add_run(part[2:-2])
            r.bold = True
        else:
            p.add_run(part)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(3)

def add_diagram_image(image_filename, caption_text):
    img_path = os.path.join(diagrams_dir, image_filename)
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(4)
        run = p_img.add_run()
        run.add_picture(img_path, width=Inches(5.8))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_cap = p_cap.add_run(caption_text)
        r_cap.font.name = 'Times New Roman'
        r_cap.font.size = Pt(10)
        r_cap.font.italic = True
        p_cap.paragraph_format.space_after = Pt(10)

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

# HALAMAN SAMPUL
add_title("LAPORAN AKHIR CAPSTONE DESIGN & PROJECT\n\nPERANCANGAN SISTEM OTOMASI PEMANTAUAN KEBENCANAAN PUSDALOPS BNPB BERBASIS INTELIJEN TEKS DAN DUAL OUTPUT GENERATOR\n\n")

p_meta = doc.add_paragraph()
p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_meta = p_meta.add_run("Diajukan untuk Memenuhi Syarat Kelulusan Mata Kuliah Capstone Design\nProgram Studi S1 Sistem Informasi\nFakultas Rekayasa Industri\nUniversitas Telkom\n\n2026\n\n")
r_meta.font.name = "Times New Roman"
r_meta.font.size = Pt(12)

doc.add_page_break()

# PROSES SEMUA FILE BAB
bab_files = [
    ('BAB_1_PENDAHULUAN.md', [
        ('gambar_1_1_gap_analisis.png', 'Gambar 1.1 Grafik Perbandingan Kesenjangan Kinerja Operasional'),
        ('gambar_1_2_fishbone_diagram.png', 'Gambar 1.2 Diagram Tulang Ikan (Fishbone Diagram) Analisis Akar Masalah')
    ]),
    ('BAB_2_TINJAUAN_PUSTAKA.md', []),
    ('BAB_3_SPESIFIKASI_PERANCANGAN.md', [
        ('gambar_3_1_sistematika_perancangan.png', 'Gambar 3.1 Diagram Alir Sistematika Perancangan Sistem Otomasi BNPB')
    ]),
    ('BAB_4_HASIL_PERANCANGAN.md', [
        ('gambar_4_1_arsitektur_sistem.png', 'Gambar 4.1 Arsitektur Sistem 4-Tier Pemantauan Kebencanaan Pusdalops BNPB'),
        ('gambar_usecase_diagram.png', 'Gambar 4.2 Use Case Diagram Sistem Otomasi Pusdalops BNPB (Adel)'),
        ('gambar_erd_database.png', 'Gambar 4.3 Entity Relationship Diagram (ERD) Basis Data Relasional (Seto)')
    ]),
    ('BAB_5_ANALISIS_BIAYA_KELAYAKAN.md', [
        ('gambar_5_1_analisis_kelayakan_bep.png', 'Gambar 5.1 Grafik Titik Impas (Break-Even Point) Sistem Otomasi Pusdalops BNPB')
    ]),
    ('BAB_6_EVALUASI_VALIDASI.md', []),
    ('BAB_7_KESIMPULAN_SARAN.md', []),
    ('LAMPIRAN_DAN_LOGBOOK.md', [])
]

for filename, diagrams in bab_files:
    filepath = os.path.join(docs_dir, filename)
    if not os.path.exists(filepath):
        continue
    
    with open(filepath, 'r', encoding='utf-8') as f:
        lines = f.readlines()
        
    in_mermaid = False
    in_code = False
    in_table = False
    table_lines = []

    for line in lines:
        stripped = line.strip()
        
        # Tangani blok mermaid / code block
        if stripped.startswith('```mermaid'):
            in_mermaid = True
            continue
        elif in_mermaid and stripped.startswith('```'):
            in_mermaid = False
            continue
        elif in_mermaid:
            continue
            
        if stripped.startswith('```python') or stripped.startswith('```text'):
            in_code = True
            continue
        elif in_code and stripped.startswith('```'):
            in_code = False
            continue
        elif in_code:
            p_code = doc.add_paragraph()
            r_c = p_code.add_run(stripped)
            r_c.font.name = 'Consolas'
            r_c.font.size = Pt(9.5)
            p_code.paragraph_format.line_spacing = 1.0
            p_code.paragraph_format.space_after = Pt(2)
            continue

        # Lewati penyematan markdown gambar jika ada (kita masukkan via add_diagram_image)
        if stripped.startswith('![') or (stripped.startswith('*Gambar') and stripped.endswith('*')):
            continue

        # Tangani tabel markdown
        if stripped.startswith('|') and stripped.endswith('|'):
            table_lines.append(stripped)
            in_table = True
            continue
        else:
            if in_table and table_lines:
                # Render tabel yang terkumpul
                rows_data = []
                for tline in table_lines:
                    if '---' in tline:
                        continue
                    cols = [c.strip() for c in tline.split('|')[1:-1]]
                    if any(cols):
                        rows_data.append(cols)
                if rows_data:
                    n_rows = len(rows_data)
                    n_cols = len(rows_data[0])
                    table = doc.add_table(rows=n_rows, cols=n_cols)
                    table.alignment = WD_TABLE_ALIGNMENT.CENTER
                    table.autofit = True
                    for r_idx, row in enumerate(rows_data):
                        for c_idx, val in enumerate(row):
                            if c_idx < n_cols:
                                cell = table.cell(r_idx, c_idx)
                                cell.text = val
                                for cp in cell.paragraphs:
                                    cp.paragraph_format.line_spacing = 1.15
                                    cp.paragraph_format.space_after = Pt(2)
                                    for crun in cp.runs:
                                        crun.font.name = 'Times New Roman'
                                        crun.font.size = Pt(9.5)
                                        if r_idx == 0:
                                            crun.bold = True
                                if r_idx == 0:
                                    set_cell_background(cell, "EAECEE")
                    doc.add_paragraph().paragraph_format.space_after = Pt(4)
                table_lines = []
                in_table = False

        if not stripped:
            continue
            
        # Heading 1, 2, 3
        if stripped.startswith('# '):
            add_heading_1(stripped[2:])
        elif stripped.startswith('## '):
            add_heading_2(stripped[3:])
        elif stripped.startswith('### '):
            add_heading_3(stripped[4:])
        elif stripped.startswith('* ') or stripped.startswith('- '):
            add_bullet(stripped[2:])
        else:
            add_body_paragraph(stripped)

    # Masukkan gambar diagram yang relevan pada akhir bab jika ada
    for img_file, cap_text in diagrams:
        add_diagram_image(img_file, cap_text)

    doc.add_page_break()

doc.save(output_docx)
print(f"Laporan Capstone Word (.docx) berhasil dibuat di: {output_docx}")
